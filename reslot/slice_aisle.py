#!/usr/bin/env python3
"""
Single-aisle pilot view of the floor-2 plan (run after run_all.py).
Pulls the racks, category assignment and bin-to-bin moves that touch one aisle
out of the full-floor plan, so the aisle can be re-slotted on its own first.

Usage:
    python3 slice_aisle.py --aisle K --out path/to/K_Aisle_Reslot.xlsx
"""
import argparse, os, pathlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
HERE=pathlib.Path(__file__).parent

ap=argparse.ArgumentParser()
ap.add_argument('--aisle',required=True)
ap.add_argument('--out',required=True)
a=ap.parse_args(); A=a.aisle.upper()

inv=pd.read_pickle(HERE/'packed.pkl'); rt=pd.read_pickle(HERE/'racks_assigned.pkl'); bd=pd.read_pickle(HERE/'bins.pkl')
binlab=bd.set_index('binCode').label.to_dict(); binaisle=bd.set_index('binCode').aisle.to_dict()
inv['From Rack']=inv.bin_code.map(binlab)
inv['from_in']=inv.bin_code.map(binaisle)==A
inv['to_in']=inv['To Bin'].map(binaisle)==A
sl=inv[inv.from_in|inv.to_in].copy()
sl['Flow']=sl.apply(lambda r:'Stays (no move)' if r['Move?']=='No' else
    'Within aisle' if r.from_in and r.to_in else 'Into aisle' if r.to_in else 'Out of aisle',axis=1)

racks=rt[rt.aisle==A].sort_values('seq')
tin=sl[sl.to_in]
# one row per unit (category / dedicated subcategory) with its consecutive rack run on the whole floor
allr=rt.sort_values('seq'); u2r={}
for r_ in allr.itertuples():
    for u in str(r_.category).split(' + '): u2r.setdefault(u,[]).append(r_.label)
units=list(dict.fromkeys(u for c in racks.category for u in str(c).split(' + ')))
cnt=tin.groupby('To Category').agg(SKUs=('sku_id','nunique'),Units=('live_inv_qty','sum'),Moves_In=('Move?',lambda s:(s=='Yes').sum()))
cat=pd.DataFrame([dict(u=u,Zone=racks[racks.category.str.contains(u,regex=False)].zone.iloc[0],
    Racks_in_aisle=', '.join(l for l in u2r[u] if l.startswith(A+'-')),Racks_on_floor=', '.join(u2r[u]),
    Only_in_this_aisle='Yes' if all(l.startswith(A+'-') for l in u2r[u]) else 'No (continues '+', '.join(l for l in u2r[u] if not l.startswith(A+'-'))+')') for u in units])
cat=cat.join(cnt,on='u').fillna({'SKUs':0,'Units':0,'Moves_In':0})
cat=cat.rename(columns={'u':'Unit (Category / Subcategory)','Racks_in_aisle':'Racks in '+A,'Racks_on_floor':'All racks on floor 2','Only_in_this_aisle':'Only in this aisle?','Moves_In':'Moves In'})

# rack layout with slot fill
bk=bd[(bd.aisle==A)&(bd.cls!='OVERHEAD')]
bka=bd[bd.aisle==A]
slots=bk.groupby('label').lim.sum(); capv=bka.groupby('label').binvol.sum()
tin=tin.assign(v=tin.live_inv_qty*tin.sku_vol_cc)
used=tin.drop_duplicates(['sku_id','To Bin']).groupby('To Rack').size(); nsku=tin.groupby('To Rack').sku_id.nunique()
uv=tin.groupby('To Rack').v.sum()
ra=racks[['label','domtier','category','zone']].copy()
ra['Slots']=ra.label.map(slots).fillna(0).astype(int); ra['SKUs After']=ra.label.map(nsku).fillna(0).astype(int)
ra['Slot Fill %']=(ra.label.map(used).fillna(0)/ra.Slots.where(ra.Slots>0)*100).round(0).fillna(0).astype(int)
ra['Stock Vol L']=(ra.label.map(uv).fillna(0)/1000).round(0).astype(int)
ra['Vol Fill % (excl. overhead)']=(ra.label.map(tin[tin['To Bin Class']!='OVERHEAD'].groupby('To Rack').v.sum()).fillna(0)/ra.label.map(bk.groupby('label').binvol.sum())*100).round(0).fillna(0).astype(int)
ra.columns=['Rack','Bin Tier','Category','Zone','Slots','SKUs After','Slot Fill %','Stock Vol L','Vol Fill % (excl. overhead)']

flow=sl.groupby('Flow').agg(Rows=('sku_id','size'),SKUs=('sku_id','nunique'),Units=('live_inv_qty','sum')).reset_index()
outd=sl[sl.Flow=='Out of aisle'].assign(Dest=lambda d:d['To Rack'].str.split('-').str[0]).groupby(['Dest','To Category']).size().reset_index(name='Moves').sort_values('Moves',ascending=False)
outd.columns=['To Aisle','To Category','Moves']

mv=sl.rename(columns={'sku_id':'SKU ID','product_name':'SKU Name','category_name':'Old Category','To Category':'New Category',
  'sku_vol_cc':'SKU Vol cc','size':'Size','bin_code':'From Bin','live_inv_qty':'Qty to Move'})
mv=mv[['Flow','SKU ID','SKU Name','Old Category','New Category','SKU Vol cc','Size','From Bin','From Rack','To Bin','To Rack','To Bin Class','Qty to Move','Move Vol cc','Rule']]
order={'Out of aisle':0,'Within aisle':1,'Into aisle':2,'Stays (no move)':3}
mv=mv.sort_values(['Flow','To Rack','To Bin','SKU Name'],key=lambda s:s.map(order) if s.name=='Flow' else s)
mv['SKU Vol cc']=mv['SKU Vol cc'].fillna(0).round(0).astype(int)

# checks on the aisle
lim=bd.set_index('binCode').lim.to_dict(); bv=bd.set_index('binCode').binvol.to_dict(); cls=bd.set_index('binCode').cls.to_dict()
occ=tin.drop_duplicates(['sku_id','To Bin']).groupby('To Bin').size()
sk=pd.read_pickle(HERE/'sku.pkl').set_index('sku_id')
parent_of=lambda c:str(c).split(' :: ')[0]
FOODP={p_ for p_ in sk.unit.str.split(' :: ').str[0].unique() if p_ in set(x for c in rt[rt.zone=='FOOD'].parent for x in str(c).split(' + '))}
occv=tin.groupby('To Bin').v.sum(); tcls=bd.set_index('binCode').binTypeCode.to_dict()
par=tin.sku_id.map(sk.unit).str.split(' :: ').str[0]; rpar=tin['To Rack'].map(rt.set_index('label').parent)
rcat=tin['To Rack'].map(rt.set_index('label').category)
offcat=tin[[not (str(u) in str(c).split(' + ')) for u,c in zip(tin['To Category'],rcat)]]
rpar=rpar.map(lambda c:str(c).split(' + ')[0])
restr=tin.sku_id.map(sk.liquid)|par.isin(['Premium Beauty','Paan Corner'])
chk=[('Moves from/to outside floor 2 (Dry_2)',int((~sl.bin_code.isin(bd.binCode)).sum()+(~sl['To Bin'].isin(bd.binCode)).sum())),
('Comingle breaches (distinct SKUs > bin limit)',int(sum(c>lim.get(b,99) for b,c in occ.items()))),
('Bin volume breaches (qty x unit vol > btc_vol_cc, excl. D_BAT)',int(sum(v>bv[b]+1 for b,v in occv.items() if tcls.get(b)!='D_BAT_1'))),
('Unit bigger than bin (excl. D_BAT)',int(((tin['To Bin'].map(bv)<tin.sku_vol_cc)&(tin['To Bin'].map(tcls)!='D_BAT_1')).sum())),
('Liquid/beauty/Paan in overhead or reserve',int((restr&tin['To Bin'].map(cls).isin(['OVERHEAD','RESERVE'])).sum())),
('Food SKU on non-food rack or vice versa',int((par.isin(FOODP)!=rpar.map(lambda p_:parent_of(p_) in FOODP)).sum())),
('SKUs of this aisle\'s categories placed outside their rack run',int(inv[inv['To Category'].isin(units)&(inv['To Bin']!='OVERFLOW')].pipe(lambda d:d[[r_ not in u2r.get(u,[]) for u,r_ in zip(d['To Category'],d['To Rack'])]]).sku_id.nunique())),
('SKUs with no bin (overflow)',int(inv[inv['To Bin']=='OVERFLOW'].sku_id.nunique())),
('SKUs on a rack not assigned to their category',int(offcat.sku_id.nunique()))]

HDR=PatternFill('solid',fgColor='1F3864'); HF=Font(name='Arial',bold=True,color='FFFFFF',size=10)
CF=Font(name='Arial',size=10); TF=Font(name='Arial',bold=True,size=13,color='1F3864'); H2=Font(name='Arial',bold=True,size=11,color='1F3864')
th=Side(style='thin',color='D9D9D9'); BORD=Border(th,th,th,th)
CEN=Alignment('center','center'); LEF=Alignment('left','center',wrap_text=True); ALT=PatternFill('solid',fgColor='F2F5FA')
FOODF=PatternFill('solid',fgColor='E2EFDA'); NFOODF=PatternFill('solid',fgColor='FCE4D6')
FLOWF={'Out of aisle':'FCE4D6','Within aisle':'FFF2CC','Into aisle':'E2EFDA','Stays (no move)':'EDEDED'}
wb=Workbook()
def table(ws,df,r0,widths=None):
    for j,c in enumerate(df.columns,1):
        x=ws.cell(r0,j,c); x.fill=HDR; x.font=HF; x.alignment=CEN; x.border=BORD
    for r,(_,row) in enumerate(df.iterrows(),r0+1):
        for j,c in enumerate(df.columns,1):
            v=row[c]; v=None if (isinstance(v,float) and pd.isna(v)) else (int(v) if hasattr(v,'item') and float(v).is_integer() else v)
            x=ws.cell(r,j,v); x.font=CF; x.border=BORD; x.alignment=LEF
            if (r-r0)%2==0: x.fill=ALT
    for col,w in (widths or {}).items(): ws.column_dimensions[col].width=w
    return r0+len(df)+2
def sheet(name,df,title,widths):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    ws['A1']=title; ws['A1'].font=TF; table(ws,df,3,widths); ws.freeze_panes='A4'; return ws

ws=wb.active; ws.title='Summary'; ws.sheet_view.showGridLines=False
ws['A1']=f'Aisle {A} Pilot Re-slot — slice of the Floor 2 (Dry_2) plan'; ws['A1'].font=TF
ws['A2']=f'Moves stay inside floor 2 (Dry_2) only · quantities in units · volume = qty x sku_vol_cc vs btc_vol_cc · {len(racks)} racks ({racks.label.iloc[0]} to {racks.label.iloc[-1]}) · {len(bd[bd.aisle==A])} bins · {int(slots.sum())} usable comingle slots'; ws['A2'].font=CF
ws['A4']='Movement'; ws['A4'].font=H2; r=table(ws,flow,5)
ws.cell(r,1,'Rule checks (aisle bins)').font=H2; r=table(ws,pd.DataFrame(chk,columns=['Check','Count']),r+1)
ws.cell(r,1,f'Where SKUs leaving {A} go').font=H2; table(ws,outd,r+1)
ws.column_dimensions['A'].width=46; ws.column_dimensions['B'].width=30; ws.column_dimensions['C'].width=10

w2=sheet('Category Plan',cat[['Unit (Category / Subcategory)','Zone','Racks in '+A,'All racks on floor 2','Only in this aisle?','SKUs','Units','Moves In']],f'Aisle {A} — category / subcategory -> consecutive racks (a category lives only in its run)',
  {'A':40,'B':11,'C':22,'D':30,'E':26,'F':7,'G':7,'H':9})
for i in range(4,4+len(cat)): w2.cell(i,2).fill=FOODF if w2.cell(i,2).value=='FOOD' else NFOODF
w3=sheet('Rack Layout',ra,f'Aisle {A} — rack -> category (one category per rack; a boundary rack may be shared by the two adjacent categories of the same food/non-food side)',{'A':8,'B':9,'C':60,'D':11,'E':7,'F':10,'G':10,'H':10,'I':12})
for i in range(4,4+len(ra)):
    f=FOODF if w3.cell(i,4).value=='FOOD' else NFOODF
    for j in range(1,10): w3.cell(i,j).fill=f
w4=sheet('Bin-to-Bin Moves',mv,f'Aisle {A} — every SKU-bin leaving, moving within, or entering the aisle',
  {'A':15,'B':30,'C':42,'D':18,'E':24,'F':9,'G':7,'H':14,'I':8,'J':14,'K':8,'L':10,'M':9,'N':10,'O':40})
for i in range(4,4+len(mv)): w4.cell(i,1).fill=PatternFill('solid',fgColor=FLOWF.get(w4.cell(i,1).value,'FFFFFF'))
os.makedirs(os.path.dirname(os.path.abspath(a.out)),exist_ok=True); wb.save(a.out)
print(ra.to_string(index=False)); print(flow.to_string(index=False)); print(chk); print('saved',a.out)
